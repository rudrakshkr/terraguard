"""Build HillSense_Engineering_Day.pptx — 10 slides, 16:9, clean design system."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette
INK      = RGBColor(0x0F, 0x17, 0x2A)
SLATE    = RGBColor(0x33, 0x41, 0x55)
MUTED    = RGBColor(0x64, 0x74, 0x8B)
FAINT    = RGBColor(0x94, 0xA3, 0xB8)
LIGHT    = RGBColor(0xF8, 0xFA, 0xFC)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
BORDER   = RGBColor(0xE2, 0xE8, 0xF0)
GREEN    = RGBColor(0x15, 0x80, 0x3D)
GREEN_BG = RGBColor(0xDC, 0xFC, 0xE7)
AMBER    = RGBColor(0xB4, 0x53, 0x09)
AMBER_BG = RGBColor(0xFE, 0xF3, 0xC7)
RED      = RGBColor(0xB9, 0x1C, 0x1C)
RED_BG   = RGBColor(0xFE, 0xE2, 0xE2)
BLUE     = RGBColor(0x1D, 0x4E, 0xD8)
BLUE_BG  = RGBColor(0xDB, 0xEA, 0xFE)
M1, M2, M3 = RGBColor(0x1B, 0x27, 0x40), RGBColor(0x22, 0x30, 0x52), RGBColor(0x18, 0x23, 0x3B)

SW, SH = 13.333, 7.5
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- helpers
def rect(slide, x, y, w, h, fill, line=None, radius=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try: shape.adjustments[0] = radius
        except Exception: pass
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    if line: shape.line.color.rgb = line; shape.line.width = Pt(1)
    else: shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def tri(slide, x, y, w, h, fill):
    sp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp

def oval(slide, x, y, d, fill):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp

def para(tf, segs, align=PP_ALIGN.LEFT, before=0, after=6, first=False, line=None):
    """segs: list of (text, size, color, bold) — one paragraph, mixed runs."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_before = Pt(before); p.space_after = Pt(after)
    if line: p.line_spacing = line
    for t, size, color, bold in segs:
        r = p.add_run(); r.text = t
        f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = FONT
    return p

def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def chip(slide, x, y, text, fg, bg, size=11.5, bold=True, border=None):
    w = 0.32 + len(text) * size * 0.0075
    c = rect(slide, x, y, w, 0.34, bg, line=border, radius=0.5)
    tf = c.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = fg; r.font.name = FONT
    return w

def new_slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, SW, SH, bg)
    return s

def header(slide, kicker, title, num):
    tf = textbox(slide, 0.6, 0.42, 12.1, 0.32)
    para(tf, [(kicker.upper(), 12.5, GREEN, True)], first=True, after=0)
    tf = textbox(slide, 0.6, 0.72, 12.1, 0.62)
    para(tf, [(title, 28, INK, True)], first=True, after=0)
    rect(slide, 0.6, 1.52, 12.13, 0.016, BORDER)
    foot(slide, num)

def foot(slide, num):
    tf = textbox(slide, 0.6, 7.08, 6.0, 0.3)
    para(tf, [("HillSense — Engineering Day", 10, MUTED, False)], first=True, after=0)
    tf = textbox(slide, 11.7, 7.08, 1.03, 0.3)
    para(tf, [(f"{num:02d}", 10, MUTED, False)], align=PP_ALIGN.RIGHT, first=True, after=0)

def card(slide, x, y, w, h, title, body, accent=GREEN, tsize=15, bsize=12.5, badge=None):
    rect(slide, x, y, w, h, WHITE, line=BORDER, radius=0.07)
    rect(slide, x, y + 0.18, 0.055, h - 0.36, accent)
    ty = y + 0.22
    if badge:
        chip(slide, x + 0.25, ty, badge, accent, {GREEN: GREEN_BG, AMBER: AMBER_BG, RED: RED_BG, BLUE: BLUE_BG}[accent], size=10.5)
        ty += 0.5
    tf = textbox(slide, x + 0.25, ty, w - 0.45, h - (ty - y) - 0.15)
    para(tf, [(title, tsize, INK, True)], first=True, after=4)
    para(tf, [(body, bsize, SLATE, False)], line=1.08, after=0)

# ================================================================ 1 · TITLE
s = new_slide(INK)
tri(s, -1.2, 5.1, 5.2, 2.4, M1)
tri(s, 2.4, 4.5, 4.6, 3.0, M2)
tri(s, 6.1, 5.3, 4.4, 2.2, M3)
tri(s, 9.2, 4.8, 4.13, 2.7, M1)
tri(s, 11.4, 5.6, 1.93, 1.9, M2)
rect(s, 0, 7.32, SW, 0.18, GREEN)

tf = textbox(s, 1.667, 1.30, 10.0, 1.1)
para(tf, [("ENGINEERING DAY", 52, WHITE, True)], align=PP_ALIGN.CENTER, first=True, after=0)
rect(s, (SW - 2.2) / 2, 2.52, 2.2, 0.05, GREEN)
tf = textbox(s, 1.667, 2.86, 10.0, 0.85)
para(tf, [("HillSense", 40, GREEN, True)], align=PP_ALIGN.CENTER, first=True, after=0)
tf = textbox(s, 1.667, 3.78, 10.0, 0.5)
para(tf, [("AI-verified hazard alerts for mountain communities", 17, FAINT, False)],
     align=PP_ALIGN.CENTER, first=True, after=0)
tf = textbox(s, 1.667, 5.02, 10.0, 1.0)
para(tf, [("Made By: Rudraksh, Jatin", 20, RGBColor(0xE2, 0xE8, 0xF0), True)],
     align=PP_ALIGN.CENTER, first=True, after=4)
para(tf, [("B.Tech CSE 5th Sem", 15, FAINT, False)], align=PP_ALIGN.CENTER, after=0)

# ================================================================ 2 · PROBLEM
s = new_slide()
header(s, "The problem", "Mountain hazards, no reliable signal", 2)
tf = textbox(s, 0.6, 1.95, 6.9, 4.6)
bullets = [
    ("Every monsoon, landslides, rockfalls and flash floods cut off roads across Himachal — often with no warning.", True),
    ("Official advisories arrive slowly; travellers depend on WhatsApp forwards and word-of-mouth.", True),
    ("Rumours spread panic, while genuine, life-saving information gets buried in the noise.", True),
    ("Worse — once a hazard clears, stale warnings keep circulating for days.", True),
]
first = True
for text, _ in bullets:
    para(tf, [("•  ", 15.5, GREEN, True), (text, 15.5, SLATE, False)],
         first=first, after=14, line=1.15)
    first = False
tf = textbox(s, 0.6, 6.1, 6.9, 0.6)
para(tf, [("The real gap: ", 15, INK, True),
          ("between when a hazard happens and when people hear about it — verified.", 15, SLATE, False)],
     first=True, after=0, line=1.1)
rect(s, 7.9, 2.25, 4.83, 3.3, WHITE, line=BORDER, radius=0.06)
rect(s, 7.9, 2.25, 4.83, 0.09, GREEN)
tf = textbox(s, 8.25, 2.62, 4.13, 2.1)
para(tf, [("“Is the road I'm about to take actually safe right now?”", 21, INK, True)],
     first=True, after=8, line=1.12)
para(tf, [("— every traveller, trucker and local in the hills", 13, MUTED, False)], after=0)

# ================================================================ 3 · SOLUTION
s = new_slide()
header(s, "The solution", "HillSense — crowdsourced alerts with an AI evidence gate", 3)
tf = textbox(s, 0.6, 1.9, 12.13, 0.75)
para(tf, [("A community reporting platform where every public alert must survive an AI verification check before it reaches anyone. ",
           16.5, SLATE, False), ("Report → Verify → Alert → Sustain.", 16.5, INK, True)],
     first=True, after=0, line=1.15)
pillars = [
    ("1 · REPORT", "Structured, geo-tagged reports with photos — hazard type, timing, severity, casualties.", GREEN),
    ("2 · VERIFY", "An AI evidence gate cross-checks image, text, location and nearby reports.", BLUE),
    ("3 · ALERT", "Only evidence-consistent incidents enter the public nearby feed and map, badged AI CHECK PASSED.", AMBER),
    ("4 · SUSTAIN", "The community confirms “still present” or “cleared”; incident status stays tied to current community updates.", GREEN),
]
x = 0.6
for title, body, accent in pillars:
    card(s, x, 3.0, 2.9, 2.35, title, body, accent=accent, tsize=14, bsize=12)
    x += 3.08
rect(s, 0.6, 5.75, 12.13, 0.85, GREEN_BG, radius=0.12)
tf = textbox(s, 0.95, 5.98, 11.5, 0.45, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Publication rule:  ", 13.5, GREEN, True),
          ("only evidence-consistent reports are eligible for the public feed — NEEDS REVIEW is held for review, REJECTED is not published.",
           13.5, SLATE, False)], first=True, after=0)

# ================================================================ 4 · HOW IT WORKS
s = new_slide()
header(s, "How it works", "From a report on the hillside to a trusted incident update", 4)
steps = [
    ("Report", "Text + photo + structured context, pinned to the reporter's actual GPS location."),
    ("Verify", "Evidence checks compare image, text, location and related reports; human review handles ambiguity."),
    ("Publish", "Evidence-consistent reports are eligible for the public feed; unclear reports stay under review."),
    ("Confirm", "One response per user: “still present” or “has cleared” — server-enforced."),
    ("Update status", "Community updates help keep the incident current and reduce stale information."),
]
x, y, w, gap = 0.6, 2.35, 2.2, 0.28
for i, (title, body) in enumerate(steps, 1):
    rect(s, x, y, w, 3.15, WHITE, line=BORDER, radius=0.09)
    oval(s, x + 0.25, y + 0.25, 0.42, GREEN)
    tf = textbox(s, x + 0.25, y + 0.295, 0.42, 0.34)
    para(tf, [(str(i), 15, WHITE, True)], align=PP_ALIGN.CENTER, first=True, after=0)
    tf = textbox(s, x + 0.25, y + 0.85, w - 0.5, 2.1)
    para(tf, [(title, 15.5, INK, True)], first=True, after=5)
    para(tf, [(body, 11.8, SLATE, False)], line=1.1, after=0)
    if i < 5:
        tf = textbox(s, x + w, y + 1.32, gap, 0.5)
        para(tf, [("→", 20, GREEN, True)], align=PP_ALIGN.CENTER, first=True, after=0)
    x += w + gap
tf = textbox(s, 0.6, 5.95, 12.13, 0.6)
para(tf, [("Evidence in, rumours out — ", 14, INK, True),
          ("the pipeline judges the report, never the reporter.", 14, SLATE, False)],
     align=PP_ALIGN.CENTER, first=True, after=0)

# ================================================================ 5 · AI VERIFICATION
s = new_slide()
header(s, "The core engine", "AI verification — it judges evidence, not people", 5)
checks = [
    "Image ↔ description consistency",
    "Image ↔ hazard-type consistency",
    "Text coherence & specificity",
    "A plausible hazard actually visible",
    "Location & timestamp plausibility",
    "Duplicate / recycled media detection",
    "Nearby corroborating reports",
    "Image quality & evidence sufficiency",
    "Contradiction flags & clear-report votes",
]
tf = textbox(s, 0.6, 1.95, 6.6, 0.4)
para(tf, [("Every report runs the same evidence checklist:", 14.5, SLATE, True)], first=True, after=0)
for i, c in enumerate(checks):
    col, row = divmod(i, 5)
    tf = textbox(s, 0.6 + col * 3.55, 2.5 + row * 0.62, 3.4, 0.55)
    para(tf, [(f"{i+1:02d}  ", 13, GREEN, True), (c, 13, SLATE, False)], first=True, after=0, line=1.05)
outcomes = [
    ("✓ VERIFIED", "Consistent available evidence — eligible for the public feed, badged AI CHECK PASSED.", GREEN, GREEN_BG),
    ("⚠ NEEDS REVIEW", "Evidence insufficient or ambiguous — held in a review queue, not published as fact.", AMBER, AMBER_BG),
    ("✕ REJECTED", "Contradicted or implausible — e.g. “highway buried” over an ordinary road photo. Never published.", RED, RED_BG),
]
y = 2.0
for title, body, accent, bg in outcomes:
    rect(s, 7.75, y, 4.98, 1.42, bg, radius=0.09)
    tf = textbox(s, 8.05, y + 0.18, 4.4, 1.1)
    para(tf, [(title, 15.5, accent, True)], first=True, after=4)
    para(tf, [(body, 12, SLATE, False)], line=1.08, after=0)
    y += 1.6
tf = textbox(s, 7.75, 6.75, 4.98, 0.4)
para(tf, [("Corroboration matters:", 12.5, INK, True),
          (" multiple nearby reports add independent evidence.", 12.5, SLATE, False)], first=True, after=0)

# ================================================================ 6 · TRUST LAYER
s = new_slide()
header(s, "Trust & safety", "A community layer built for honesty", 6)
cards = [
    ("Phone + OTP sign-in", "No passwords anywhere — hashed one-time codes, expiring bearer sessions, real SMS provider ready.", BLUE),
    ("One confirm per user", "user + incident + response is enforced server-side; refreshing the page never unlocks a second vote.", GREEN),
    ("Private by default", "Comments show first names only. Phone numbers are never displayed or stored in public records.", BLUE),
    ("Comments ≠ verified facts", "Community observations are visually distinct from the AI CHECK PASSED assessment — and labelled as such.", AMBER),
    ("Corroboration, not reputation", "Nearby reports reinforce an incident's evidence. No user trust scores — people aren't ranked.", GREEN),
    ("Honest geolocation", "Real GPS → reverse-geocoded full address, editable before saving. If it fails: clearly marked “approximate”.", GREEN),
]
for i, (title, body, accent) in enumerate(cards):
    col, row = i % 3, i // 3
    card(s, 0.6 + col * 4.11, 1.95 + row * 2.42, 3.9, 2.2, title, body, accent=accent, tsize=14.5, bsize=12)

# ================================================================ 7 · FEATURES
s = new_slide()
header(s, "The product", "Everything a hill traveller actually needs", 7)
cards = [
    ("Structured report form", "Hazard type, timing, “happening now?”, what's affected, people trapped, observed severity — concise, progressive.", GREEN),
    ("Personalized nearby feed", "Verified incidents ranked by distance, severity, confirmations and recency — signed in after onboarding.", BLUE),
    ("List + map views", "Browse alerts as clean cards or as pins on a live map of your radius.", GREEN),
    ("Confirmations & comments", "“Still present / has cleared” with one vote per user, plus threaded community observations.", AMBER),
    ("Corroboration chips", "Multi-report incidents display “Corroborated by N related reports” right on the card.", BLUE),
    ("Graceful fallbacks", "Manual location entry, heuristic verification without an LLM key, honest approx-labelling throughout.", GREEN),
]
for i, (title, body, accent) in enumerate(cards):
    col, row = i % 3, i // 3
    card(s, 0.6 + col * 4.11, 1.95 + row * 2.42, 3.9, 2.2, title, body, accent=accent, tsize=14.5, bsize=12)

# ================================================================ 8 · TECH STACK
s = new_slide()
header(s, "Under the hood", "A modern, deployable stack", 8)
rect(s, 0.6, 1.95, 5.95, 3.5, WHITE, line=BORDER, radius=0.07)
rect(s, 0.6, 1.95, 5.95, 0.07, BLUE)
tf = textbox(s, 0.95, 2.25, 5.3, 3.1)
para(tf, [("FRONTEND", 12.5, BLUE, True)], first=True, after=8)
for line in ["Next.js App Router + React 19 + TypeScript",
             "Tailwind CSS design system — one shared layout grid",
             "Auth, geolocation & theme hooks on the client",
             "Verification trace & evidence UI components"]:
    para(tf, [("•  ", 13.5, GREEN, True), (line, 13.5, SLATE, False)], after=7, line=1.1)
rect(s, 6.78, 1.95, 5.95, 3.5, WHITE, line=BORDER, radius=0.07)
rect(s, 6.78, 1.95, 5.95, 0.07, GREEN)
tf = textbox(s, 7.13, 2.25, 5.3, 3.1)
para(tf, [("BACKEND & AI", 12.5, GREEN, True)], first=True, after=8)
for line in ["Next.js API routes (Node runtime)",
             "OTP auth — hashed codes & hashed sessions",
             "LLM vision verification + heuristic fallback pipeline",
             "OpenStreetMap Nominatim reverse geocoding",
             "File-based stores with mtime cache invalidation"]:
    para(tf, [("•  ", 13.5, GREEN, True), (line, 13.5, SLATE, False)], after=7, line=1.1)
rect(s, 0.6, 5.75, 12.13, 0.8, RGBColor(0xF1, 0xF5, 0xF9), radius=0.12)
tf = textbox(s, 0.6, 5.98, 12.13, 0.4, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Browser  →  /api/auth · /api/incidents · /api/analyze · /api/geo  →  stores · LLM · Nominatim",
           13.5, SLATE, False)], align=PP_ALIGN.CENTER, first=True, after=0)

# ================================================================ 9 · DEMO
s = new_slide()
header(s, "Live demo", "Watch a rumour become a verified alert", 9)
demo = [
    ("Sign in", "Phone + OTP — no passwords. (Dev mode shows the code on screen, honestly labelled.)"),
    ("Set location", "One tap on “Use my location” → real GPS → full reverse-geocoded address, editable before saving."),
    ("Report a hazard", "Photo + description + structured context: type, severity, people affected."),
    ("Watch verification", "Live trace — image analysed, description checked, corroboration found — then the verdict."),
    ("Go public", "The VERIFIED alert appears in the nearby feed. Confirm it, comment on it — or clear it."),
]
y = 1.95
for i, (title, body) in enumerate(demo, 1):
    oval(s, 0.6, y + 0.06, 0.46, GREEN)
    tf = textbox(s, 0.6, y + 0.115, 0.46, 0.36)
    para(tf, [(str(i), 15, WHITE, True)], align=PP_ALIGN.CENTER, first=True, after=0)
    tf = textbox(s, 1.35, y, 11.3, 0.8)
    para(tf, [(title + "   ", 15.5, INK, True), (body, 14, SLATE, False)], first=True, after=0, line=1.1)
    y += 0.88
rect(s, 0.6, 6.45, 12.13, 0.55, AMBER_BG, radius=0.16)
tf = textbox(s, 0.9, 6.57, 11.6, 0.35, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Demo backup ready:  ", 12.5, AMBER, True),
          ("screenshots and a recorded run, in case venue Wi-Fi or GPS misbehaves.", 12.5, SLATE, False)],
     first=True, after=0)

# ================================================================ 10 · CLOSE
s = new_slide()
header(s, "Wrap-up", "What HillSense changes", 10)
rect(s, 0.6, 1.95, 6.35, 3.1, WHITE, line=BORDER, radius=0.07)
rect(s, 0.6, 1.95, 6.35, 0.07, GREEN)
tf = textbox(s, 0.95, 2.25, 5.7, 2.7)
para(tf, [("IMPACT", 12.5, GREEN, True)], first=True, after=8)
for line in ["Verified information reaches travellers in minutes, not days.",
             "Rumours are filtered by evidence — and the community keeps alerts current.",
             "One private place that answers: “Is this route safe right now?”"]:
    para(tf, [("•  ", 13.5, GREEN, True), (line, 13.5, SLATE, False)], after=8, line=1.12)
rect(s, 7.18, 1.95, 5.55, 3.1, WHITE, line=BORDER, radius=0.07)
rect(s, 7.18, 1.95, 5.55, 0.07, BLUE)
tf = textbox(s, 7.53, 2.25, 4.9, 0.4)
para(tf, [("ROADMAP", 12.5, BLUE, True)], first=True, after=0)
road = ["Real SMS delivery (Twilio / MSG91)", "Admin review queue for NEEDS REVIEW",
        "SMS & push alerts", "Offline-first PWA", "Govt. disaster-management feed integration"]
y = 2.68
for r in road:
    w = chip(s, 7.53, y, r, SLATE, RGBColor(0xF1, 0xF5, 0xF9), size=12, bold=False, border=BORDER)
    y += 0.46
tf = textbox(s, 0.6, 5.55, 12.13, 1.2)
para(tf, [("Thank you.", 34, INK, True)], align=PP_ALIGN.CENTER, first=True, after=4)
para(tf, [("Questions?   ·   Rudraksh & Jatin — B.Tech CSE 5th Sem", 14.5, MUTED, False)],
     align=PP_ALIGN.CENTER, after=0)

# ---------------------------------------------------------------- save
prs.core_properties.title = "Engineering Day — HillSense"
prs.core_properties.author = "Rudraksh, Jatin"
out = "HillSense_Engineering_Day.pptx"
prs.save(out)
print(f"Saved {out} with {len(prs.slides.slides._sldIdLst)} slides" if False else f"Saved {out}")